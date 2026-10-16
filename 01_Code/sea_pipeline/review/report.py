from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .models import ReviewOutput
from .tags import FAIL_TAG_COLORS, choose_fail_tag, fail_tag_for_criterion, normalize_fail_tag


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
FAIL_FILL = PatternFill("solid", fgColor="F4CCCC")
PASS_FILL = PatternFill("solid", fgColor="D9EAD3")


def _write_rows(ws: Worksheet, rows: list[dict[str, Any]]) -> None:
    if not rows:
        ws.append(["message"])
        ws.append(["No rows produced."])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    for column_cells in ws.columns:
        max_len = min(max(len(str(cell.value or "")) for cell in column_cells), 80)
        ws.column_dimensions[column_cells[0].column_letter].width = max(max_len + 2, 12)
        header = str(column_cells[0].value or "")
        if header == "summary":
            ws.column_dimensions[column_cells[0].column_letter].width = 120
        elif header == "value":
            ws.column_dimensions[column_cells[0].column_letter].width = 100
        for cell in column_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if str(cell.value).lower() == "fail":
                cell.fill = FAIL_FILL
            elif str(cell.value).lower() == "pass":
                cell.fill = PASS_FILL
    if "fail_tag" in headers:
        tag_column = headers.index("fail_tag") + 1
        for cell in ws.iter_cols(min_col=tag_column, max_col=tag_column, min_row=2):
            for tag_cell in cell:
                tag = str(tag_cell.value or "").strip()
                color = FAIL_TAG_COLORS.get(tag)
                if color:
                    tag_cell.fill = PatternFill("solid", fgColor=color)
                    tag_cell.font = Font(bold=True)


SUMMARY_COMPONENTS = ("takeaway", "concept", "resource", "quiz_question", "poll", "case_study")
SUMMARY_LABELS = {
    "takeaway": "takeaways",
    "concept": "concepts",
    "resource": "resources",
    "quiz_question": "quiz",
    "poll": "polls",
    "case_study": "case_studies",
}
SUMMARY_COMPONENT_ALIASES = {
    "takeaways": "takeaway",
    "concepts": "concept",
    "resources": "resource",
    "quiz": "quiz_question",
    "polls": "poll",
    "case_studies": "case_study",
}
DETAIL_COLUMNS_WITH_STATUS = [
    "lesson_id",
    "module_id",
    "component",
    "component_index",
    "overall_status",
    "word_count",
    "missing",
    "context_accurate_ai_review",
    "duplicates",
    "fail_tag",
    "summary",
    "value",
]
DETAIL_COLUMNS_FAIL_ONLY = [
    column for column in DETAIL_COLUMNS_WITH_STATUS
    if column != "overall_status"
]
SPELLCHECK_COLUMNS = ["lesson_id", "frame", "misspelled_word", "sentence"]


def _module_sort_key(module_id: str) -> tuple[int, str]:
    try:
        return int(module_id), module_id
    except ValueError:
        return 9999, module_id


def _detail_component_key(component: str) -> str:
    return SUMMARY_COMPONENT_ALIASES.get(component, component)


def _criterion_category(criterion: str) -> str:
    normalized = criterion.lower()
    if "llm" in normalized or "semantic" in normalized:
        return "context_accurate_ai_review"
    if "word count" in normalized:
        return "word_count"
    if "duplicate" in normalized or "unique" in normalized:
        return "duplicates"
    if "blob" in normalized:
        return "missing"
    if normalized.startswith("has ") or "missing expected" in normalized or "present" in normalized:
        return "missing"
    return "_other_deterministic"


def _has_llm_rows(output: ReviewOutput) -> bool:
    return any(
        _criterion_category(str(row.get("criterion", ""))) == "context_accurate_ai_review"
        for row in output.component_rows
    )


def _append_note(existing: str, note: str) -> str:
    note = str(note or "").strip()
    if not note:
        return existing
    if not existing:
        return note
    return f"{existing} | {note}"


def _new_aggregated_detail(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "lesson_id": row.get("lesson_id", ""),
        "module_id": row.get("module_id", ""),
        "component": _detail_component_key(str(row.get("component", ""))),
        "component_index": row.get("component_index", ""),
        "overall_status": "pass",
        "word_count": "pass",
        "missing": "pass",
        "context_accurate_ai_review": "pass",
        "duplicates": "pass",
        "_other_deterministic": "pass",
        "fail_tag": "",
        "summary": "",
        "value": row.get("value", ""),
        "source_json": row.get("source_json", ""),
    }


def _merge_detail_row(target: dict[str, Any], row: dict[str, Any]) -> None:
    category = _criterion_category(str(row.get("criterion", "")))
    status = str(row.get("status", "")).lower()
    if status == "fail":
        target[category] = "fail"
        target["overall_status"] = "fail"
        notes = str(row.get("notes") or "").strip()
        suggested_fix = str(row.get("suggested_fix") or "").strip()
        criterion = str(row.get("criterion") or "").strip()
        fail_tag = normalize_fail_tag(row.get("fail_tag"), default=fail_tag_for_criterion(criterion))
        target["fail_tag"] = choose_fail_tag(target.get("fail_tag"), fail_tag)
        target["summary"] = _append_note(target["summary"], notes or criterion)
        if suggested_fix and suggested_fix != notes:
            target["summary"] = _append_note(target["summary"], f"Fix: {suggested_fix}")
    row_value = str(row.get("value") or "")
    target_value = str(target.get("value") or "")
    if row_value and len(row_value) > len(target_value):
        target["value"] = row.get("value")


def build_summary_rows(output: ReviewOutput, *, include_ai_review: bool | None = None) -> list[dict[str, Any]]:
    fail_counts: dict[str, Counter[str]] = {}
    for row in build_component_detail_rows(output, keep_pass=False, include_ai_review=include_ai_review):
        module_id = str(row.get("module_id", ""))
        component = SUMMARY_COMPONENT_ALIASES.get(str(row.get("component", "")), str(row.get("component", "")))
        if component not in SUMMARY_COMPONENTS:
            continue
        fail_counts.setdefault(module_id, Counter())[component] += 1
    for row in output.gap_rows:
        module_id = str(row.get("module_id", ""))
        component = SUMMARY_COMPONENT_ALIASES.get(str(row.get("component", "")), str(row.get("component", "")))
        if component not in SUMMARY_COMPONENTS:
            continue
        fail_counts.setdefault(module_id, Counter())[component] += 1

    module_ids = sorted(
        {str(row.get("module_id", "")) for row in output.lesson_rows if str(row.get("module_id", ""))},
        key=_module_sort_key,
    )
    rows: list[dict[str, Any]] = []
    totals = Counter()
    for module_id in module_ids:
        counts = fail_counts.get(module_id, Counter())
        row: dict[str, Any] = {"module": module_id}
        module_total = 0
        for component in SUMMARY_COMPONENTS:
            value = counts.get(component, 0)
            row[SUMMARY_LABELS[component]] = value
            module_total += value
            totals[component] += value
        row["total"] = module_total
        rows.append(row)

    total_row: dict[str, Any] = {"module": "Total"}
    grand_total = 0
    for component in SUMMARY_COMPONENTS:
        value = totals.get(component, 0)
        total_row[SUMMARY_LABELS[component]] = value
        grand_total += value
    total_row["total"] = grand_total
    rows.append(total_row)
    return rows


def build_component_detail_rows(
    output: ReviewOutput,
    *,
    keep_pass: bool,
    include_ai_review: bool | None = None,
) -> list[dict[str, Any]]:
    aggregated: dict[tuple[str, str, str], dict[str, Any]] = {}
    for row in output.component_rows:
        component = _detail_component_key(str(row.get("component", "")))
        key = (str(row.get("lesson_id", "")), component, str(row.get("component_index", "")))
        if key not in aggregated:
            aggregated[key] = _new_aggregated_detail(row)
        _merge_detail_row(aggregated[key], row)
    for row in output.gap_rows:
        component = _detail_component_key(str(row.get("component", "")))
        key = (str(row.get("lesson_id", "")), component, "missing")
        if key not in aggregated:
            aggregated[key] = {
                "lesson_id": row.get("lesson_id", ""),
                "module_id": row.get("module_id", ""),
                "component": component,
                "component_index": "missing",
                "overall_status": "fail",
                "word_count": "pass",
                "missing": "fail",
                "context_accurate_ai_review": "pass",
                "duplicates": "pass",
                "_other_deterministic": "pass",
                "fail_tag": "missing",
                "summary": "",
                "value": "",
                "source_json": row.get("source_json", ""),
            }
        aggregated[key]["summary"] = _append_note(
            aggregated[key]["summary"],
            f"{row.get('notes', '')} Expected {row.get('expected', '')}; found {row.get('actual', '')}.",
        )
    rows = [row for row in aggregated.values() if keep_pass or row["overall_status"] == "fail"]
    rows.sort(key=lambda row: (str(row["module_id"]), str(row["lesson_id"]), str(row["component"]), str(row["component_index"])))
    include_ai_review = _has_llm_rows(output) if include_ai_review is None else include_ai_review
    columns = DETAIL_COLUMNS_WITH_STATUS if keep_pass else DETAIL_COLUMNS_FAIL_ONLY
    if not include_ai_review:
        columns = [column for column in columns if column != "context_accurate_ai_review"]
    return [{column: row.get(column, "") for column in columns} for row in rows]


def load_spellcheck_rows(spellcheck_csv_path: str | Path | None) -> list[dict[str, Any]]:
    if spellcheck_csv_path is None:
        return []
    path = Path(spellcheck_csv_path)
    if not path.exists():
        return []
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = []
        for row in reader:
            rows.append({column: row.get(column, "") for column in SPELLCHECK_COLUMNS})
    return rows


def write_review_workbook(
    output: ReviewOutput,
    output_path: str | Path,
    *,
    keep_pass: bool = False,
    include_ai_review: bool | None = None,
    spellcheck_csv_path: str | Path | None = None,
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    output.summary_rows = build_summary_rows(output, include_ai_review=include_ai_review)
    component_detail_rows = build_component_detail_rows(
        output,
        keep_pass=keep_pass,
        include_ai_review=include_ai_review,
    )
    spellcheck_rows = load_spellcheck_rows(spellcheck_csv_path)
    spellcheck_path_exists = spellcheck_csv_path is not None and Path(spellcheck_csv_path).exists()
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [
        ("Summary", output.summary_rows),
        ("Component Details", component_detail_rows),
    ]
    if spellcheck_path_exists:
        sheets.append(("Spellcheck", spellcheck_rows or [{column: "" for column in SPELLCHECK_COLUMNS}]))
    for name, rows in sheets:
        ws = workbook.create_sheet(name)
        _write_rows(ws, rows)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    workbook.save(path)
    output.output_path = path
    return path


def _bar(percent: float, width: int = 24) -> str:
    filled = round((percent / 100) * width)
    return "█" * filled + "░" * (width - filled)


def review_completion_summary(workbook_path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Summary" not in workbook.sheetnames:
        raise ValueError(f"Review workbook does not contain a 'Summary' sheet: {workbook_path}")
    sheet = workbook["Summary"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "") for value in rows[0]]
    total_row = next((row for row in rows[1:] if str(row[0]) == "Total"), None)
    if total_row is None:
        return []
    totals: list[dict[str, Any]] = []
    for column_index, header in enumerate(headers[1:], start=1):
        if header == "total":
            continue
        failed = int(total_row[column_index] or 0)
        percent = 100.0 if failed == 0 else 0.0
        totals.append(
            {
                "category": header,
                "passed": 0,
                "failed": failed,
                "total": failed,
                "percent": percent,
                "bar": _bar(percent),
            }
        )
    return totals


def format_review_completion_summary(workbook_path: str | Path) -> str:
    rows = review_completion_summary(workbook_path)
    if not rows:
        return "No review completion rows found."
    label_width = max(len(str(row["category"])) for row in rows)
    lines = [f"Review completion summary: {workbook_path}"]
    for row in rows:
        label = str(row["category"]).replace("_", " ").title().ljust(label_width)
        lines.append(f"{label} |{row['bar']}| {row['failed']} fail(s)")
    return "\n".join(lines)


def print_review_completion_summary(workbook_path: str | Path) -> None:
    print(format_review_completion_summary(workbook_path))
