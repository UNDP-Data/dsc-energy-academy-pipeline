from __future__ import annotations

from pathlib import Path

from docx import Document


def read_docx_text(path: str | Path) -> str:
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def write_docx_text(text: str, path: str | Path) -> Path:
    output_path = Path(path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for line in (text or "").splitlines():
        if line.strip():
            document.add_paragraph(line)
    document.save(output_path)
    return output_path
