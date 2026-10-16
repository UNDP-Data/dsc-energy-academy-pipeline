from __future__ import annotations

import csv
import hashlib
import html
import json
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from sea_pipeline.shared.config import PipelineConfig


LogFn = Callable[[str], None] | None

BR_RE = re.compile(r"(?i)<br\s*/?>")
AI_TAG_RE = re.compile(r"(?is)<ai\b[^>]*>(.*?)</ai>")
A_TAG_RE = re.compile(r"(?is)<a\b[^>]*href=[\"']([^\"']+)[\"'][^>]*>(.*?)</a>")
OTHER_TAG_RE = re.compile(r"<[^>]+>")
HTML_BOLD_RE = re.compile(r"(?is)<(strong|b)>(.*?)</\1>")
HTML_ITALIC_RE = re.compile(r"(?is)<(em|i)>(.*?)</\1>")
MD_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
MD_ITALIC_RE = re.compile(r"(?<!\w)_(.+?)_(?!\w)")
WHITESPACE_RE = re.compile(r"[ \t]+")
URL_RE = re.compile(r"https?://[^\s)>\"]+")

DOCX_OK_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
MEDIA_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
IMAGE_WIDTH = Inches(6.3)
TEXT_TEMPLATE_STYLES: dict[str, dict[str, object]] = {
    "subtitle": {"size": 16, "bold": True},
    "subtitle_small": {"size": 13, "bold": True},
    "paragraph_large": {"size": 13},
    "paragraph_medium": {"size": 11},
    "paragraph_small": {"size": 9},
    "bullet_point": {"size": 11},
    "bullet_point_with_number": {"size": 11},
    "bullet_point_with_highlight": {"size": 11},
    "kpi_highlight_large": {"size": 18, "bold": True},
    "kpi_highlight_medium": {"size": 15, "bold": True},
    "quote_large_with_name": {"size": 15, "italic": True},
    "quote_large_without_name": {"size": 15, "italic": True},
    "quote_small_with_name": {"size": 12, "italic": True},
    "quote_small_without_name": {"size": 12, "italic": True},
}

UI_TEXT_KEYS = {
    "cta",
    "button_cta",
    "buttonCta",
    "nextButton",
    "submitButton",
    "retryButton",
    "startButton",
}
NON_CONTENT_KEYS = {
    "template_id",
    "color_scheme",
    "type",
    "progress",
    "id",
    "lessonId",
    "nextLessonId",
    "nextBlockId",
    "pollId",
    "quizId",
    "src",
    "poster",
    "href",
    "url",
    "image",
    "frame",
}
TEXT_KEYS = {
    "label",
    "title",
    "subtitle",
    "intro",
    "text",
    "body",
    "description",
    "value",
    "prompt",
    "question",
    "quote",
    "author",
    "source",
    "footnote",
    "correct",
    "wrong",
    "result",
    "score",
    "passingScore",
    "passed",
    "failed",
}
NAVIGATION_TEMPLATES = {"connection_next", "connection_back"}
MEDIA_PLACEHOLDER_TEMPLATES = {"embed", "video", "video_full_height", "lesson_cover_video"}
CHART_TEMPLATES = {"chart", "echarts_chart"}
PHOTO_TEMPLATES = {
    "photo_full_height",
    "photo_vertical",
    "photo-horizontal",
    "lesson_cover",
    "lesson_part_cover",
    "lesson_subpart_cover",
    "module_cover",
    "chapter_cover",
}
GENERIC_TEXT_TEMPLATES = {"module_cover", "chapter_cover", "lesson_cover", "lesson_part_cover", "lesson_subpart_cover", "chapter_outro", "module_outro", "lesson_outro"}
WARNING_AUDIT_TYPES = {
    "missing_image_src",
    "image_download_failed",
    "image_conversion_failed",
    "image_insert_failed",
    "missing_lesson_json",
    "invalid_lesson_json",
    "unsupported_template",
}


@dataclass
class ExportSettings:
    lesson_root: Path
    structure_path: Path
    output_dir: Path
    image_cache_dir: Path
    audit_path: Path
    language: str = "en"
    include_images: bool = True
    include_placeholders: bool = True
    module_ids: set[str] | None = None


@dataclass
class ExportAuditEntry:
    module_id: str
    lesson_id: str
    template_id: str
    issue_type: str
    detail: str
    reference: str = ""


@dataclass
class ExportContext:
    settings: ExportSettings
    audit: list[ExportAuditEntry] = field(default_factory=list)
    lessons_written: int = 0
    images_inserted: int = 0
    placeholders_inserted: int = 0
    missing_lessons: list[str] = field(default_factory=list)

    def add_audit(self, module_id: str, lesson_id: str, template_id: str, issue_type: str, detail: str, reference: str = "") -> None:
        self.audit.append(
            ExportAuditEntry(
                module_id=module_id,
                lesson_id=lesson_id,
                template_id=template_id,
                issue_type=issue_type,
                detail=detail,
                reference=reference,
            )
        )


def _natural_sort_key(value: object) -> list[object]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", str(value))]


def _safe_filename(value: str, maxlen: int = 90) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9 _.-]+", "", value or "").strip()
    cleaned = re.sub(r"\s+", " ", cleaned)
    return (cleaned[:maxlen].rstrip() or "Module").strip()


def _load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _clean_html_text(value: Any, *, keep_links: bool = True) -> str:
    text = "" if value is None else str(value)
    if not text:
        return ""
    text = html.unescape(text).replace("\u00a0", " ")
    text = BR_RE.sub("\n", text)
    text = AI_TAG_RE.sub(r"\1", text)

    def replace_link(match: re.Match[str]) -> str:
        href = html.unescape(match.group(1)).strip()
        inner = _clean_html_text(match.group(2), keep_links=False)
        if keep_links and href and inner and href != inner:
            return f"{inner} ({href})"
        return inner or href

    text = A_TAG_RE.sub(replace_link, text)
    lines = [WHITESPACE_RE.sub(" ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _strip_all_tags(value: Any) -> str:
    text = _clean_html_text(value)
    text = OTHER_TAG_RE.sub("", text)
    lines = [WHITESPACE_RE.sub(" ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _inline_text(value: Any) -> str:
    text = OTHER_TAG_RE.sub("", "" if value is None else str(value))
    return WHITESPACE_RE.sub(" ", text)


def _parse_inline_runs(value: Any) -> list[tuple[str, dict[str, bool]]]:
    text = _clean_html_text(value)

    def apply_markdown(plain: str, style: dict[str, bool]) -> list[tuple[str, dict[str, bool]]]:
        parts: list[tuple[str, str]] = []
        output: list[tuple[str, dict[str, bool]]] = []
        last = 0
        for match in MD_BOLD_RE.finditer(plain):
            if match.start() > last:
                parts.append(("plain", plain[last:match.start()]))
            parts.append(("bold", match.group(1)))
            last = match.end()
        if last < len(plain):
            parts.append(("plain", plain[last:]))
        for kind, segment in parts:
            if kind == "bold":
                next_style = dict(style)
                next_style["bold"] = True
                output.append((segment, next_style))
                continue
            last_italic = 0
            for match in MD_ITALIC_RE.finditer(segment):
                if match.start() > last_italic:
                    output.append((segment[last_italic:match.start()], dict(style)))
                next_style = dict(style)
                next_style["italic"] = True
                output.append((match.group(1), next_style))
                last_italic = match.end()
            if last_italic < len(segment):
                output.append((segment[last_italic:], dict(style)))
        return [(chunk, chunk_style) for chunk, chunk_style in output if _inline_text(chunk)]

    def walk_html(fragment: str, inherited: dict[str, bool]) -> list[tuple[str, dict[str, bool]]]:
        output: list[tuple[str, dict[str, bool]]] = []
        while fragment:
            bold_match = HTML_BOLD_RE.search(fragment)
            italic_match = HTML_ITALIC_RE.search(fragment)
            if not bold_match and not italic_match:
                output.extend(apply_markdown(fragment, inherited))
                break
            if bold_match and italic_match:
                match, style_key = (bold_match, "bold") if bold_match.start() < italic_match.start() else (italic_match, "italic")
            elif bold_match:
                match, style_key = bold_match, "bold"
            else:
                match, style_key = italic_match, "italic"
            if match.start() > 0:
                output.extend(apply_markdown(fragment[:match.start()], inherited))
            next_style = dict(inherited)
            next_style[style_key] = True
            output.extend(walk_html(match.group(2), next_style))
            fragment = fragment[match.end():]
        return output

    runs = []
    for chunk, style in walk_html(text, {"bold": False, "italic": False}):
        cleaned = _inline_text(chunk)
        if cleaned:
            runs.append((cleaned, style))
    return runs


def _new_document() -> Document:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    return document


def _apply_run_format(run: Any, *, font_size: int | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    if font_size is not None:
        run.font.size = Pt(font_size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _add_rich_paragraph(
    document: Document,
    text: Any,
    *,
    style: str | None = None,
    bold_all: bool = False,
    italic_all: bool = False,
    font_size: int | None = None,
):
    cleaned = _clean_html_text(text)
    if not cleaned:
        return None
    paragraph = document.add_paragraph(style=style) if style else document.add_paragraph()
    if bold_all or italic_all:
        run = paragraph.add_run(_strip_all_tags(cleaned))
        _apply_run_format(run, font_size=font_size, bold=bold_all, italic=italic_all)
        return paragraph
    for chunk, run_style in _parse_inline_runs(cleaned):
        run = paragraph.add_run(chunk)
        _apply_run_format(
            run,
            font_size=font_size,
            bold=bool(run_style.get("bold")),
            italic=bool(run_style.get("italic")),
        )
    return paragraph


def _add_multiline(
    document: Document,
    text: Any,
    *,
    style: str | None = None,
    bold_all: bool = False,
    italic_all: bool = False,
    font_size: int | None = None,
) -> None:
    for line in _clean_html_text(text).splitlines():
        if line.strip():
            _add_rich_paragraph(document, line, style=style, bold_all=bold_all, italic_all=italic_all, font_size=font_size)


def _add_bullet(document: Document, text: Any, *, font_size: int | None = None) -> None:
    _add_rich_paragraph(document, text, style="List Bullet", font_size=font_size)


def _add_numbered(document: Document, text: Any, *, font_size: int | None = None) -> None:
    _add_rich_paragraph(document, text, style="List Number", font_size=font_size)


def _add_placeholder(document: Document, text: str, context: ExportContext) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.italic = True
    context.placeholders_inserted += 1


def _text_value(value: Any) -> str:
    return _strip_all_tags(value)


def _content(segment: dict[str, Any]) -> dict[str, Any]:
    content = segment.get("content")
    return content if isinstance(content, dict) else {}


def _image_src(value: Any) -> str:
    if isinstance(value, dict):
        src = value.get("src")
        return str(src).strip() if src else ""
    return ""


def _looks_like_image_url(url: str) -> bool:
    parsed = urlparse(url)
    suffix = Path(unquote(parsed.path)).suffix.lower()
    return parsed.scheme in {"http", "https", "file"} and suffix in MEDIA_IMAGE_EXTS


def _media_display_path(value: Any) -> str:
    raw = _text_value(value)
    if not raw:
        return ""
    parsed = urlparse(raw)
    if parsed.scheme in {"http", "https", "file"} and parsed.path:
        path = unquote(parsed.path).lstrip("/")
        if path.startswith("images/"):
            return path.removeprefix("images/")
        return path or raw
    return raw


def _cache_path_for_url(url: str, cache_dir: Path) -> Path:
    parsed = urlparse(url)
    suffix = Path(unquote(parsed.path)).suffix.lower() or ".bin"
    digest = hashlib.sha256(url.encode("utf-8")).hexdigest()[:24]
    return cache_dir / f"{digest}{suffix}"


def _download_or_resolve_image(url: str, cache_dir: Path) -> Path | None:
    parsed = urlparse(url)
    if parsed.scheme == "file":
        path = Path(unquote(parsed.path))
        return path if path.exists() else None
    if parsed.scheme == "" and Path(url).exists():
        return Path(url)
    if parsed.scheme not in {"http", "https"}:
        return None
    cache_dir.mkdir(parents=True, exist_ok=True)
    output = _cache_path_for_url(url, cache_dir)
    if output.exists() and output.stat().st_size > 0:
        return output
    try:
        import requests

        response = requests.get(url, timeout=60)
        response.raise_for_status()
        output.write_bytes(response.content)
        return output
    except Exception:
        return None


def _docx_compatible_image(path: Path) -> Path | None:
    if not path.exists() or path.stat().st_size == 0:
        return None
    if path.suffix.lower() in DOCX_OK_EXTS:
        return path
    try:
        from PIL import Image

        image = Image.open(path)
        image.load()
        if image.mode not in {"RGB", "RGBA"}:
            image = image.convert("RGB")
        output = path.with_suffix(".png")
        image.save(output, format="PNG")
        return output
    except Exception:
        return None


def _insert_image(document: Document, src: str, caption: str, context: ExportContext, *, module_id: str, lesson_id: str, template_id: str) -> bool:
    if not src:
        context.add_audit(module_id, lesson_id, template_id, "missing_image_src", "Image source is blank.")
        return False
    image_path = _download_or_resolve_image(src, context.settings.image_cache_dir)
    if image_path is None:
        context.add_audit(module_id, lesson_id, template_id, "image_download_failed", "Could not fetch image.", src)
        return False
    compatible = _docx_compatible_image(image_path)
    if compatible is None:
        context.add_audit(module_id, lesson_id, template_id, "image_conversion_failed", "Could not convert image for DOCX.", str(image_path))
        return False
    try:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(compatible), width=IMAGE_WIDTH)
        if caption:
            caption_paragraph = document.add_paragraph()
            caption_run = caption_paragraph.add_run(caption)
            caption_run.italic = True
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        context.images_inserted += 1
        return True
    except Exception as exc:
        context.add_audit(module_id, lesson_id, template_id, "image_insert_failed", str(exc), str(compatible))
        return False


def _render_text_fields(document: Document, obj: Any, *, skip_keys: set[str] | None = None, prefix: str = "") -> None:
    skip = skip_keys or set()
    if isinstance(obj, dict):
        for key, value in obj.items():
            if key in skip or key in UI_TEXT_KEYS or key in NON_CONTENT_KEYS:
                continue
            if key in TEXT_KEYS and isinstance(value, str):
                label = prefix or key.replace("_", " ").title()
                if key in {"title", "label"}:
                    _add_multiline(document, value, bold_all=True)
                elif key in {"subtitle", "source", "footnote", "author"}:
                    _add_multiline(document, f"{label}: {_text_value(value)}", italic_all=True)
                else:
                    _add_multiline(document, value)
            elif isinstance(value, (dict, list)):
                _render_text_fields(document, value, skip_keys=skip, prefix=prefix)
    elif isinstance(obj, list):
        for item in obj:
            _render_text_fields(document, item, skip_keys=skip, prefix=prefix)


def _render_object_list(document: Document, title: str, items: list[Any], fields: list[tuple[str, str]]) -> None:
    if title:
        _add_rich_paragraph(document, title, bold_all=True)
    for index, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            continue
        parts = []
        for key, label in fields:
            value = _text_value(item.get(key, ""))
            if value:
                parts.append(f"{label}: {value}" if label else value)
        if parts:
            _add_bullet(document, f"{index}. " + " | ".join(parts))


def _render_key_concepts(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    _add_multiline(document, content.get("title", "Key Concepts"), bold_all=True)
    for concept in content.get("concepts", []) if isinstance(content.get("concepts"), list) else []:
        if not isinstance(concept, dict):
            continue
        title = _text_value(concept.get("title", ""))
        body = _text_value(concept.get("body", ""))
        source = _text_value(concept.get("source", ""))
        if title and body:
            _add_bullet(document, f"{title}: {body}")
        elif title or body:
            _add_bullet(document, title or body)
        if source:
            _add_multiline(document, f"Source: {source}", italic_all=True)


def _render_text_segment(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    text_elements = content.get("text_elements")
    if not isinstance(text_elements, list):
        _render_generic(document, segment)
        return
    for element in text_elements:
        if not isinstance(element, dict):
            continue
        element_template = str(element.get("template_id", "") or "").strip()
        element_content = element.get("content") if isinstance(element.get("content"), dict) else {}
        text = element_content.get("text", "")
        if not _text_value(text):
            continue
        template_style = TEXT_TEMPLATE_STYLES.get(element_template, {})
        font_size = int(template_style["size"]) if "size" in template_style else None
        bold = bool(template_style.get("bold", False))
        italic = bool(template_style.get("italic", False))
        if "bullet" in element_template:
            _add_bullet(document, text, font_size=font_size)
        elif "number" in element_template:
            _add_numbered(document, text, font_size=font_size)
        elif bold or italic or font_size:
            _add_multiline(document, text, bold_all=bold, italic_all=italic, font_size=font_size)
        else:
            _add_multiline(document, text)


def _render_key_takeaways(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    _add_multiline(document, content.get("title", "Key Takeaways"), bold_all=True)
    for item in content.get("takeaways", []) if isinstance(content.get("takeaways"), list) else []:
        if not isinstance(item, dict):
            continue
        title = _text_value(item.get("title", ""))
        description = _text_value(item.get("description", ""))
        text = _text_value(item.get("text", ""))
        combined = text or " ".join(part for part in [title, description] if part)
        if combined:
            _add_bullet(document, combined)


def _render_learning_objectives(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    _add_multiline(document, content.get("title", "Learning Objectives"), bold_all=True)
    if content.get("intro"):
        _add_multiline(document, content.get("intro"))
    for objective in content.get("objectives", []) if isinstance(content.get("objectives"), list) else []:
        if not isinstance(objective, dict):
            continue
        title = _text_value(objective.get("title", ""))
        description = _text_value(objective.get("description", ""))
        combined = " ".join(part for part in [title, description] if part)
        if combined:
            _add_bullet(document, combined)


def _render_list_of_lessons(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    _add_multiline(document, content.get("title", "Lessons"), bold_all=True)
    lessons = content.get("lessons", []) if isinstance(content.get("lessons"), list) else []
    for lesson in lessons:
        if not isinstance(lesson, dict):
            continue
        title = _text_value(lesson.get("title", ""))
        lesson_id = _text_value(lesson.get("lessonId", ""))
        description = _text_value(lesson.get("description", ""))
        parts = [part for part in [lesson_id, title, description] if part]
        if parts:
            _add_bullet(document, " | ".join(parts))


def _render_key_resources(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    _add_multiline(document, content.get("title", "Key Resources"), bold_all=True)
    for resource in content.get("resources", []) if isinstance(content.get("resources"), list) else []:
        if not isinstance(resource, dict):
            continue
        title = _text_value(resource.get("text", "") or resource.get("title", ""))
        href = str(resource.get("href", "") or "").strip()
        line = title
        if href:
            line = f"{line} ({href})" if line else href
        if line:
            _add_bullet(document, line)


def _render_infographic(document: Document, segment: dict[str, Any], context: ExportContext, *, module_id: str, lesson_id: str) -> None:
    content = _content(segment)
    metadata = content.get("metadata") if isinstance(content.get("metadata"), dict) else {}
    image = content.get("image") if isinstance(content.get("image"), dict) else {}
    title = _text_value(metadata.get("title", "Infographic"))
    _add_rich_paragraph(document, f"Infographic: {title}" if title else "Infographic", bold_all=True)
    src = _image_src(image)
    inserted = False
    if context.settings.include_images:
        inserted = _insert_image(document, src, _text_value(image.get("caption", "")), context, module_id=module_id, lesson_id=lesson_id, template_id="infographic")
    if not inserted and context.settings.include_placeholders:
        _add_placeholder(document, f"[INFOGRAPHIC IMAGE: {src or 'missing source'}]", context)
    for key in ("subtitle", "description", "source", "footnote"):
        value = _text_value(metadata.get(key, ""))
        if value:
            label = key.replace("_", " ").title()
            _add_multiline(document, f"{label}: {value}", italic_all=key in {"source", "footnote"})


def _render_chart_placeholder(document: Document, segment: dict[str, Any], context: ExportContext) -> None:
    content = _content(segment)
    metadata = content.get("metadata") if isinstance(content.get("metadata"), dict) else {}
    chart_id = _text_value(content.get("chart_id", "") or content.get("id", "") or metadata.get("id", ""))
    title = _text_value(content.get("title", "") or metadata.get("title", "") or chart_id)
    label = " | ".join(part for part in [chart_id, title] if part) or "unknown chart"
    if context.settings.include_placeholders:
        _add_placeholder(document, f"[CHART PLACEHOLDER: {label}]", context)
    _render_text_fields(document, metadata)
    _render_text_fields(document, content, skip_keys={"metadata", "chart", "data", "options"})


def _render_embed_placeholder(document: Document, segment: dict[str, Any], context: ExportContext) -> None:
    content = _content(segment)
    title = _text_value(content.get("title", ""))
    url = _text_value(content.get("url", ""))
    label = " | ".join(part for part in [title, url] if part) or "missing embed URL"
    if context.settings.include_placeholders:
        _add_placeholder(document, f"[EMBED PLACEHOLDER: {label}]", context)
    if title:
        _add_multiline(document, title, bold_all=True)
    if url:
        _add_multiline(document, f"Embed URL: {url}")


def _render_image_hotspot_placeholder(document: Document, segment: dict[str, Any], context: ExportContext) -> None:
    content = _content(segment)
    image = content.get("image") if isinstance(content.get("image"), dict) else {}
    hotspots = content.get("hotspots", []) if isinstance(content.get("hotspots"), list) else []
    src = _image_src(image)
    if context.settings.include_placeholders:
        _add_placeholder(document, f"[IMAGE HOTSPOT PLACEHOLDER: {src or 'missing image'} | hotspots: {len(hotspots)}]", context)
    for index, hotspot in enumerate(hotspots, start=1):
        if isinstance(hotspot, dict):
            label = _text_value(hotspot.get("label", "") or hotspot.get("title", "") or hotspot.get("text", ""))
            body = _text_value(hotspot.get("body", "") or hotspot.get("description", ""))
            if label or body:
                _add_bullet(document, f"Hotspot {index}: " + " | ".join(part for part in [label, body] if part))


def _render_video_placeholder(document: Document, segment: dict[str, Any], context: ExportContext) -> None:
    content = _content(segment)
    src = _media_display_path(content.get("src", ""))
    if context.settings.include_placeholders:
        _add_placeholder(document, f"Video: {src or 'missing source'}", context)


def _render_case_study_cover(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    _add_multiline(document, content.get("intro", ""), italic_all=True)
    _add_multiline(document, content.get("title", "Case Study"), bold_all=True)
    _add_multiline(document, content.get("subtitle", ""), italic_all=True)


def _render_scored_quiz(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    _add_multiline(document, content.get("title", "Scored Quiz"), bold_all=True)
    _add_multiline(document, content.get("intro", ""))
    if content.get("passingScore") not in {None, ""}:
        _add_multiline(document, f"Passing score: {content.get('passingScore')}")
    questions = content.get("questions", []) if isinstance(content.get("questions"), list) else []
    for index, question in enumerate(questions, start=1):
        if not isinstance(question, dict):
            continue
        prompt = _text_value(question.get("prompt", ""))
        _add_numbered(document, f"Question {index}: {prompt}" if prompt else f"Question {index}")
        options = question.get("options", []) if isinstance(question.get("options"), list) else []
        for option in options:
            if not isinstance(option, dict):
                continue
            option_id = option.get("id", "")
            value = _text_value(option.get("value", ""))
            if value:
                _add_bullet(document, f"{option_id}. {value}" if option_id != "" else value)
        solution = question.get("solution")
        if solution:
            _add_multiline(document, f"Solution: {solution}", italic_all=True)
        messages = question.get("messages") if isinstance(question.get("messages"), dict) else {}
        for key in ("correct", "wrong"):
            message = messages.get(key) if isinstance(messages.get(key), dict) else {}
            body = _text_value(message.get("body", ""))
            if body:
                _add_multiline(document, f"{key.title()} feedback: {body}", italic_all=True)


def _render_poll(document: Document, segment: dict[str, Any]) -> None:
    content = _content(segment)
    _add_multiline(document, content.get("title", "Poll"), bold_all=True)
    _add_multiline(document, content.get("description", "") or content.get("prompt", ""))
    options = content.get("options", []) if isinstance(content.get("options"), list) else []
    for option in options:
        if isinstance(option, dict):
            value = _text_value(option.get("value", "") or option.get("text", ""))
            if value:
                _add_bullet(document, value)


def _render_generic(document: Document, segment: dict[str, Any]) -> None:
    _render_text_fields(document, _content(segment))


def _render_segment(document: Document, segment: dict[str, Any], context: ExportContext, *, module_id: str, lesson_id: str) -> None:
    template_id = str(segment.get("template_id", "") or "").strip()
    if template_id in NAVIGATION_TEMPLATES:
        context.add_audit(module_id, lesson_id, template_id, "skipped_navigation", "Navigation-only template skipped.")
        return
    if template_id == "text":
        _render_text_segment(document, segment)
    elif template_id == "key_concepts":
        _render_key_concepts(document, segment)
    elif template_id == "key_takeaways":
        _render_key_takeaways(document, segment)
    elif template_id == "learning_objectives":
        _render_learning_objectives(document, segment)
    elif template_id == "list_of_lessons":
        _render_list_of_lessons(document, segment)
    elif template_id == "key_resources":
        _render_key_resources(document, segment)
    elif template_id == "infographic":
        _render_infographic(document, segment, context, module_id=module_id, lesson_id=lesson_id)
    elif template_id in CHART_TEMPLATES:
        _render_chart_placeholder(document, segment, context)
    elif template_id == "embed":
        _render_embed_placeholder(document, segment, context)
    elif template_id == "image-hotspot":
        _render_image_hotspot_placeholder(document, segment, context)
    elif template_id in MEDIA_PLACEHOLDER_TEMPLATES:
        _render_video_placeholder(document, segment, context)
    elif template_id == "case_study_cover":
        _render_case_study_cover(document, segment)
    elif template_id == "scored_quiz":
        _render_scored_quiz(document, segment)
    elif template_id == "poll":
        _render_poll(document, segment)
    elif template_id in PHOTO_TEMPLATES:
        _render_generic(document, segment)
    elif template_id in GENERIC_TEXT_TEMPLATES:
        _render_generic(document, segment)
    else:
        _render_generic(document, segment)
        if template_id:
            context.add_audit(module_id, lesson_id, template_id, "unsupported_template", "Template exported with generic text renderer; add a dedicated renderer if this template needs special handling.")


def _build_lesson_index(lesson_root: Path) -> dict[str, Path]:
    index: dict[str, Path] = {}
    paths = sorted(lesson_root.glob("Module_*/*.json"), key=lambda path: _natural_sort_key(path.relative_to(lesson_root)))
    for path in paths:
        try:
            data = _load_json(path)
        except Exception:
            continue
        lesson_id = data.get("id") if isinstance(data, dict) else None
        if isinstance(lesson_id, str) and lesson_id.strip():
            index.setdefault(lesson_id.strip(), path)
    return index


def _write_audit_csv(path: Path, audit: list[ExportAuditEntry]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = ["module_id", "lesson_id", "template_id", "issue_type", "detail", "reference"]
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for entry in audit:
            writer.writerow({column: getattr(entry, column) for column in columns})
    return path


def _settings_from_config(
    config: PipelineConfig,
    *,
    with_images: bool,
    language: str | None = None,
    module_ids: Iterable[str] | None = None,
) -> ExportSettings:
    selected_language = language or str(config.stage_options.get("docx_export_language", "en"))
    lesson_root = Path(config.stage_options.get("docx_export_lesson_root", config.paths.outputs_dir / "SEA_Modules" / selected_language))
    output_dir = Path(config.stage_options.get("docx_export_output_dir", config.paths.outputs_dir / "SEA_Module_Docx"))
    image_cache_dir = Path(config.stage_options.get("docx_export_image_cache_dir", config.paths.outputs_dir / "image_temp"))
    audit_path = Path(config.stage_options.get("docx_export_audit_path", output_dir / "docx_export_audit.csv"))
    structure_path = Path(config.stage_options.get("docx_export_structure_path", lesson_root / "module_structure.json"))
    configured_modules = config.stage_options.get("docx_export_module_ids")
    ids = module_ids if module_ids is not None else configured_modules
    if isinstance(ids, str):
        ids = [part.strip() for part in ids.split(",") if part.strip()]
    module_id_set = {str(module_id) for module_id in ids} if ids else None
    return ExportSettings(
        lesson_root=lesson_root,
        structure_path=structure_path,
        output_dir=output_dir,
        image_cache_dir=image_cache_dir,
        audit_path=audit_path,
        language=selected_language,
        include_images=bool(config.stage_options.get("docx_export_include_images", with_images)),
        include_placeholders=bool(config.stage_options.get("docx_export_include_placeholders", True)),
        module_ids=module_id_set,
    )


def export_docx_materials(
    config: PipelineConfig,
    *,
    with_images: bool = True,
    language: str | None = None,
    module_ids: Iterable[str] | None = None,
    log: LogFn = print,
) -> dict[str, object]:
    started = time.time()
    settings = _settings_from_config(config, with_images=with_images, language=language, module_ids=module_ids)
    settings.output_dir.mkdir(parents=True, exist_ok=True)
    settings.image_cache_dir.mkdir(parents=True, exist_ok=True)
    structure = _load_json(settings.structure_path)
    modules = structure.get("modules") if isinstance(structure, dict) else None
    if not isinstance(modules, list):
        raise ValueError(f"Expected module_structure.json with a modules list: {settings.structure_path}")
    lesson_index = _build_lesson_index(settings.lesson_root)
    context = ExportContext(settings=settings)
    output_paths: list[str] = []

    for module in modules:
        if not isinstance(module, dict):
            continue
        module_id = str(module.get("id", "")).strip()
        if settings.module_ids is not None and module_id not in settings.module_ids:
            continue
        module_title = _text_value(module.get("title", ""))
        if log:
            log(f"[DOCX] Exporting Module {module_id}: {module_title}")
        document = _new_document()
        document.add_heading(f"Module {module_id}: {module_title}", level=0)
        for key in ("description", "estimatedTime"):
            value = _text_value(module.get(key, ""))
            if value:
                _add_multiline(document, f"{key.replace('_', ' ').title()}: {value}", italic_all=key == "estimatedTime")
        for chapter in module.get("chapters", []) if isinstance(module.get("chapters"), list) else []:
            if not isinstance(chapter, dict):
                continue
            chapter_id = str(chapter.get("id", "")).strip()
            chapter_title = _text_value(chapter.get("title", ""))
            document.add_heading(f"Chapter {chapter_id}: {chapter_title}", level=1)
            chapter_description = _text_value(chapter.get("description", ""))
            if chapter_description:
                _add_multiline(document, chapter_description, italic_all=True)
            for lesson_meta in chapter.get("lessons", []) if isinstance(chapter.get("lessons"), list) else []:
                if not isinstance(lesson_meta, dict):
                    continue
                lesson_id = str(lesson_meta.get("id", "")).strip()
                lesson_title = _text_value(lesson_meta.get("title", ""))
                lesson_path = lesson_index.get(lesson_id)
                if not lesson_path:
                    context.missing_lessons.append(lesson_id)
                    context.add_audit(module_id, lesson_id, "", "missing_lesson_json", "Lesson listed in structure but no lesson JSON was found.")
                    continue
                document.add_heading(f"Lesson {lesson_id}: {lesson_title}", level=2)
                lesson = _load_json(lesson_path)
                if not isinstance(lesson, dict):
                    context.add_audit(module_id, lesson_id, "", "invalid_lesson_json", "Lesson JSON root is not an object.", str(lesson_path))
                    continue
                segments = lesson.get("segments") if isinstance(lesson.get("segments"), list) else []
                for segment in segments:
                    if isinstance(segment, dict):
                        _render_segment(document, segment, context, module_id=module_id, lesson_id=lesson_id)
                context.lessons_written += 1
                document.add_paragraph("")
        output_path = settings.output_dir / f"M{module_id} - {_safe_filename(module_title)}.docx"
        document.save(output_path)
        output_paths.append(str(output_path))
        if log:
            log(f"[DOCX] Wrote {output_path}")

    audit_path = _write_audit_csv(settings.audit_path, context.audit)
    warning_count = sum(1 for entry in context.audit if entry.issue_type in WARNING_AUDIT_TYPES)
    return {
        "stage": "docx_export_with_images" if settings.include_images else "docx_export_text",
        "kind": "python",
        "status": "warning" if warning_count else "ok",
        "summary": {
            "language": settings.language,
            "lesson_root": str(settings.lesson_root),
            "output_dir": str(settings.output_dir),
            "modules_exported": len(output_paths),
            "lessons_written": context.lessons_written,
            "images_inserted": context.images_inserted,
            "placeholders_inserted": context.placeholders_inserted,
            "audit_entries": len(context.audit),
            "warning_entries": warning_count,
            "missing_lessons": len(set(context.missing_lessons)),
            "audit_path": str(audit_path),
        },
        "outputs": output_paths,
        "duration_seconds": time.time() - started,
    }


def run_docx_export(config: PipelineConfig, with_images: bool = False) -> dict[str, object]:
    return export_docx_materials(config, with_images=with_images)
