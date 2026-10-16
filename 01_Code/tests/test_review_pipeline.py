from __future__ import annotations

import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

CODE_DIR = Path(__file__).resolve().parents[1]
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from sea_pipeline.review.extract import classify_lesson_id, load_review_inputs
from sea_pipeline.review.blob_checks import ResourceBlobCheckResult
from sea_pipeline.review.controls import load_review_controls
from sea_pipeline.review.llm import _build_llm_review_request, _component_payload, _module_contexts, _parse_llm_json
from sea_pipeline.review.report import format_review_completion_summary, write_review_workbook
from sea_pipeline.review.runner import create_review_session, run_review_component, write_review_session_workbook
from sea_pipeline.review.rules import run_deterministic_review
from sea_pipeline.review.tags import FAIL_TAG_COLORS
from sea_pipeline.shared.config import PipelineConfig
from sea_pipeline.shared.io_json import save_json
from sea_pipeline.shared.env import EnvConfig
from sea_pipeline.shared.paths import RepoPaths
from sea_pipeline.auxiliary.spellcheck import EXTRA_SPELLCHECK_WHITELIST, strip_urls_from_text


def _normal_lesson(lesson_id: str = "1.1.1") -> dict:
    return {
        "id": lesson_id,
        "segments": [
            {
                "template_id": "key_concepts",
                "content": {
                    "concepts": [
                        {"title": "Energy access", "body": "Energy access refers to reliable and affordable availability of modern energy services for households and communities."},
                        {"title": "Energy equity", "body": "Energy equity refers to fair distribution of energy benefits, costs, and opportunities across people and places."},
                        {"title": "Clean cooking", "body": "Clean cooking refers to cooking technologies and fuels that reduce harmful household air pollution and emissions."},
                    ]
                },
            },
            {"template_id": "text", "content": {"body": "This lesson explains energy access, equity, and clean cooking in sustainable development."}},
            {
                "template_id": "poll",
                "content": {
                    "title": "Which access barrier matters most in your context?",
                    "options": [{"id": 1, "value": "Affordability"}, {"id": 2, "value": "Reliability"}],
                    "labels": {"submit": "Submit"},
                },
            },
            {
                "template_id": "key_takeaways",
                "content": {
                    "takeaways": [
                        {"title": "Energy access requires more than grid connections.", "description": "Useful access depends on affordability, reliability, safety, and the ability to support services that improve daily life."},
                        {"title": "Equity determines who benefits from energy progress.", "description": "Energy systems can reinforce inequality unless planning considers underserved households, local institutions, and vulnerable communities."},
                        {"title": "Clean cooking remains a central development priority.", "description": "Cleaner cooking options reduce health risks, free time, and support better environmental outcomes for families and communities."},
                    ]
                },
            },
            {
                "template_id": "key_resources",
                "content": {
                    "resources": [
                        {"text": "Tracking SDG7 energy progress", "href": "https://example.org/sdg7", "image": {"src": "https://example.org/1.webp"}},
                        {"text": "Energy access policy guide", "href": "https://example.org/access", "image": {"src": "https://example.org/2.webp"}},
                        {"text": "Clean cooking transition report", "href": "https://example.org/cooking", "image": {"src": "https://example.org/3.webp"}},
                    ]
                },
            },
        ],
    }


class ReviewPipelineTests(unittest.TestCase):
    def test_spellcheck_strips_urls_before_tokenization(self) -> None:
        text = (
            "Read https://medium.com/@vulcaninc/financing-the-future and "
            "https://www.sciencedirect.com/science/article/pii/S123 for more context."
        )
        stripped = strip_urls_from_text(text)
        self.assertNotIn("vulcaninc", stripped)
        self.assertNotIn("sciencedirect", stripped)
        self.assertNotIn("medium.com", stripped)

    def test_spellcheck_extra_whitelist_contains_requested_terms(self) -> None:
        expected = {
            "eda",
            "fpic",
            "ghg",
            "gis",
            "hr",
            "kv",
            "pp",
            "px",
            "yr",
            "additionality",
            "cashflows",
            "conservoltaic",
            "floatovoltaics",
            "offtake",
            "undergrid",
            "watertube",
            "woodfuel",
            "analysing",
            "centralised",
            "honour",
            "mobilise",
            "neighbouring",
            "optimisation",
            "prioritises",
            "standardised",
            "visualisation",
            "highveld",
            "jaalis",
            "watsonx",
            "buen",
            "contor",
            "convoca",
            "cooperativahidroelectrica",
            "dans",
            "desarrollo",
            "historia",
            "humano",
            "multiactor",
            "pacto",
            "poderosa",
            "renouvelables",
            "rurale",
            "sostenible",
            "vivir",
        }
        self.assertTrue(expected.issubset(EXTRA_SPELLCHECK_WHITELIST))
        self.assertFalse({"acccess", "distirbtuion", "socieities"} & EXTRA_SPELLCHECK_WHITELIST)

    def test_classifies_lesson_types(self) -> None:
        self.assertEqual(classify_lesson_id("1.1.1"), "lesson")
        self.assertEqual(classify_lesson_id("1.-1.0"), "module_quiz")
        self.assertEqual(classify_lesson_id("1.1.0"), "chapter_intro")
        self.assertEqual(classify_lesson_id("1.1.-1"), "chapter_outro")
        self.assertEqual(classify_lesson_id("1.0.0"), "module_intro_or_preface")

    def test_extracts_and_passes_complete_normal_lesson(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            lessons = load_review_inputs(root)
            self.assertEqual(len(lessons), 1)
            self.assertEqual(len(lessons[0].takeaways), 3)
            self.assertEqual(len(lessons[0].concepts), 3)
            self.assertEqual(len(lessons[0].resources), 3)
            result = run_deterministic_review(lessons)
            self.assertEqual(result.lesson_rows[0]["overall_status"], "pass")

    def test_loads_sequential_test_slice_from_start_and_count(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson("1.1.1"), module_dir / "1.1.1.json")
            save_json(_normal_lesson("1.1.2"), module_dir / "1.1.2.json")
            save_json(_normal_lesson("1.1.3"), module_dir / "1.1.3.json")
            lessons = load_review_inputs(root, start_lesson_id="1.1.1", lesson_count=2)
            self.assertEqual([lesson.meta.lesson_id for lesson in lessons], ["1.1.1", "1.1.2"])

    def test_false_lesson_count_means_full_run(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            lesson_root = root / "03_Outputs" / "SEA_Modules" / "en"
            module_dir = lesson_root / "Module_1"
            save_json(_normal_lesson("1.1.1"), module_dir / "1.1.1.json")
            save_json(_normal_lesson("1.1.2"), module_dir / "1.1.2.json")
            config = PipelineConfig(
                root_dir=root,
                paths=RepoPaths.from_root(root),
                env=EnvConfig(),
                stage_options={
                    "review_lesson_root": lesson_root,
                    "review_test_start_lesson_id": "1.1.1",
                    "review_test_lesson_count": False,
                    "review_output_path": root / "review.xlsx",
                    "review_cache_path": root / "cache.json",
                    "review_run_ai_review": False,
                    "review_check_resource_blobs": False,
                },
            )
            session = create_review_session(config, log=None)
            self.assertFalse(session.settings.test_mode)
            self.assertEqual([lesson.meta.lesson_id for lesson in session.lessons], ["1.1.1", "1.1.2"])

    def test_critical_ai_mode_enables_llm_with_critical_prompt(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            lesson_root = root / "03_Outputs" / "SEA_Modules" / "en"
            module_dir = lesson_root / "Module_1"
            save_json(_normal_lesson("1.1.1"), module_dir / "1.1.1.json")
            config = PipelineConfig(
                root_dir=root,
                paths=RepoPaths.from_root(root),
                env=EnvConfig(),
                stage_options={
                    "review_lesson_root": lesson_root,
                    "review_test_start_lesson_id": "1.1.1",
                    "review_test_lesson_count": 1,
                    "review_run_ai_review": "Critical",
                    "review_check_resource_blobs": False,
                },
            )
            session = create_review_session(config, log=None)
            self.assertTrue(session.settings.use_llm)
            self.assertEqual(session.settings.ai_mode, "critical")
            lesson = session.lessons[0]
            request = _build_llm_review_request(
                lesson,
                "concepts",
                _component_payload(lesson, "concepts"),
                session.controls,
                lesson.lesson_text,
                review_mode=session.settings.ai_mode,
            )
            self.assertEqual(request["review_mode"], "critical")
            self.assertIn("complete mismatch", request["criterion"])
            self.assertIn("fail only severe issues", request["general_guidance"].lower())

    def test_detects_missing_components_as_gaps(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"] = [segment for segment in lesson["segments"] if segment["template_id"] != "key_resources"]
            save_json(lesson, module_dir / "1.1.1.json")
            result = run_deterministic_review(load_review_inputs(root))
            self.assertEqual(result.lesson_rows[0]["resources_status"], "fail")
            self.assertEqual(result.gap_rows[0]["component"], "resources")

    def test_review_controls_override_numeric_limits(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            controls_path = root / "controls.csv"
            controls_path.write_text("section,key,value\nnormal_lesson,expected_resources,2\n", encoding="utf-8")
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            resources = lesson["segments"][-1]["content"]["resources"]
            lesson["segments"][-1]["content"]["resources"] = resources[:2]
            save_json(lesson, module_dir / "1.1.1.json")
            controls = load_review_controls(controls_path)
            result = run_deterministic_review(load_review_inputs(root), controls=controls)
            self.assertEqual(result.lesson_rows[0]["resources_status"], "pass")

    def test_missing_poll_is_not_required_by_default(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"] = [segment for segment in lesson["segments"] if segment["template_id"] != "poll"]
            save_json(lesson, module_dir / "1.1.1.json")
            result = run_deterministic_review(load_review_inputs(root))
            self.assertEqual(result.lesson_rows[0]["poll_status"], "pass")
            self.assertFalse(any(row["component"] == "poll" for row in result.gap_rows))

    def test_duplicate_concept_flags_only_later_occurrence(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][0]["content"]["concepts"][1]["title"] = "Energy access"
            save_json(lesson, module_dir / "1.1.1.json")
            result = run_deterministic_review(load_review_inputs(root))
            duplicate_rows = [
                row for row in result.component_rows
                if row["component"] == "concept" and row["criterion"] == "duplicate concept title"
            ]
            statuses = {row["component_index"]: row["status"] for row in duplicate_rows}
            self.assertEqual(statuses[1], "pass")
            self.assertEqual(statuses[2], "fail")
            self.assertEqual(statuses[3], "pass")
            self.assertIn("first appears as item 1", next(row["notes"] for row in duplicate_rows if row["component_index"] == 2))

    def test_duplicate_resource_flags_only_later_occurrence(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][4]["content"]["resources"][1]["href"] = "https://example.org/sdg7"
            save_json(lesson, module_dir / "1.1.1.json")
            result = run_deterministic_review(load_review_inputs(root))
            duplicate_rows = [
                row for row in result.component_rows
                if row["component"] == "resource" and row["criterion"] == "duplicate resource title or href"
            ]
            statuses = {row["component_index"]: row["status"] for row in duplicate_rows}
            self.assertEqual(statuses[1], "pass")
            self.assertEqual(statuses[2], "fail")
            self.assertEqual(statuses[3], "pass")
            self.assertIn("first appears as item 1", next(row["notes"] for row in duplicate_rows if row["component_index"] == 2))

    def test_resource_blob_check_flags_missing_photo(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")

            def checker(url: str) -> ResourceBlobCheckResult:
                if url.endswith("2.webp"):
                    return ResourceBlobCheckResult(status="missing", exists=False, blob_name="Resources/Module_1/2.webp", message="Blob not found")
                return ResourceBlobCheckResult(status="exists", exists=True, blob_name="Resources/Module_1/ok.webp")

            result = run_deterministic_review(load_review_inputs(root), resource_blob_checker=checker)
            self.assertEqual(result.lesson_rows[0]["resources_status"], "fail")
            missing_rows = [
                row for row in result.component_rows
                if row["criterion"] == "resource image exists in blob storage" and row["status"] == "fail"
            ]
            self.assertEqual(len(missing_rows), 1)
            workbook_path = write_review_workbook(result, root / "review.xlsx", keep_pass=False)
            workbook = load_workbook(workbook_path)
            sheet = workbook["Component Details"]
            headers = [cell.value for cell in sheet[1]]
            self.assertIn("missing", headers)
            self.assertNotIn("blob_storage", headers)
            rows = [
                {headers[index]: value for index, value in enumerate(row)}
                for row in sheet.iter_rows(min_row=2, values_only=True)
            ]
            resource_row = next(row for row in rows if row["component"] == "resource" and row["component_index"] == 2)
            self.assertEqual(resource_row["missing"], "fail")
            self.assertEqual(resource_row["fail_tag"], "photo")
            self.assertEqual(resource_row["value"], "Energy access policy guide")
            self.assertIn("Photo 2", resource_row["summary"])
            self.assertNotIn("https://example.org/2.webp", resource_row["summary"])
            self.assertNotIn("https://example.org/2.webp", resource_row["value"])

    def test_quiz_minimum_question_count(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(
                {
                    "id": "1.-1.0",
                    "segments": [
                        {
                            "template_id": "scored_quiz",
                            "content": {
                                "questions": [
                                    {
                                        "id": i,
                                        "prompt": f"Question {i}?",
                                        "options": [{"id": 1, "value": "A"}, {"id": 2, "value": "B"}, {"id": 3, "value": "C"}],
                                        "solution": [1],
                                        "messages": {"correct": {"body": "Yes"}, "wrong": {"body": "No"}},
                                    }
                                    for i in range(1, 9)
                                ]
                            },
                        }
                    ],
                },
                module_dir / "1.-1.0.json",
            )
            result = run_deterministic_review(load_review_inputs(root))
            self.assertEqual(result.lesson_rows[0]["quiz_status"], "fail")
            self.assertEqual(result.gap_rows[0]["component"], "quiz")

    def test_writes_review_workbook(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            workbook_path = write_review_workbook(output, root / "review.xlsx")
            workbook = load_workbook(workbook_path)
            self.assertIn("Summary", workbook.sheetnames)
            self.assertIn("Component Details", workbook.sheetnames)
            self.assertNotIn("Lesson Review", workbook.sheetnames)
            self.assertNotIn("Gaps", workbook.sheetnames)
            self.assertEqual(workbook["Summary"]["A1"].value, "module")
            self.assertEqual(workbook["Summary"]["B1"].value, "takeaways")

    def test_review_workbook_includes_spellcheck_csv_tab_when_present(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            spellcheck_path = root / "spellcheck_audit_report.csv"
            spellcheck_path.write_text(
                "lesson_id,frame,misspelled_word,sentence\n"
                "1.1.1,3,acccess,Energy acccess should be reliable.\n",
                encoding="utf-8",
            )
            output = run_deterministic_review(load_review_inputs(root))
            workbook_path = write_review_workbook(
                output,
                root / "review.xlsx",
                spellcheck_csv_path=spellcheck_path,
            )
            workbook = load_workbook(workbook_path)
            self.assertIn("Spellcheck", workbook.sheetnames)
            sheet = workbook["Spellcheck"]
            self.assertEqual([cell.value for cell in sheet[1]], ["lesson_id", "frame", "misspelled_word", "sentence"])
            self.assertEqual(sheet["C2"].value, "acccess")

    def test_formats_review_completion_summary_from_workbook(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            workbook_path = write_review_workbook(output, root / "review.xlsx")
            summary = format_review_completion_summary(workbook_path)
            self.assertIn("Review completion summary", summary)
            self.assertIn("Takeaways", summary)
            self.assertIn("0 fail(s)", summary)

    def test_component_details_can_drop_pass_rows_and_columns(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"] = [segment for segment in lesson["segments"] if segment["template_id"] != "key_resources"]
            save_json(lesson, module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            workbook_path = write_review_workbook(output, root / "review.xlsx", keep_pass=False)
            workbook = load_workbook(workbook_path)
            sheet = workbook["Component Details"]
            headers = [cell.value for cell in sheet[1]]
            self.assertNotIn("chapter_id", headers)
            self.assertNotIn("lesson_type", headers)
            self.assertNotIn("overall_status", headers)
            self.assertNotIn("source_json", headers)
            self.assertIn("word_count", headers)
            self.assertIn("missing", headers)
            self.assertNotIn("missing_present", headers)
            self.assertNotIn("blob_storage", headers)
            self.assertNotIn("context_accurate_ai_review", headers)
            self.assertNotIn("other_deterministic", headers)
            self.assertNotIn("structure", headers)
            self.assertIn("summary", headers)

    def test_component_details_keep_overall_status_only_when_pass_rows_are_kept(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            workbook_path = write_review_workbook(output, root / "review.xlsx", keep_pass=True)
            workbook = load_workbook(workbook_path)
            headers = [cell.value for cell in workbook["Component Details"][1]]
            self.assertIn("overall_status", headers)
            self.assertNotIn("source_json", headers)

    def test_component_details_are_one_row_per_lesson_component_index(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][3]["content"]["takeaways"][0]["title"] = ""
            lesson["segments"][3]["content"]["takeaways"][0]["description"] = "Too short."
            save_json(lesson, module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            workbook_path = write_review_workbook(output, root / "review.xlsx", keep_pass=False)
            workbook = load_workbook(workbook_path)
            sheet = workbook["Component Details"]
            headers = [cell.value for cell in sheet[1]]
            self.assertNotIn("context_accurate_ai_review", headers)
            rows = [
                {headers[index]: value for index, value in enumerate(row)}
                for row in sheet.iter_rows(min_row=2, values_only=True)
            ]
            keys = [(row["lesson_id"], row["component"], row["component_index"]) for row in rows]
            self.assertEqual(len(keys), len(set(keys)))
            takeaway_one = next(row for row in rows if row["component"] == "takeaway" and row["component_index"] == 1)
            self.assertEqual(takeaway_one["word_count"], "fail")
            self.assertEqual(takeaway_one["missing"], "pass")
            self.assertEqual(takeaway_one["fail_tag"], "wordcount")
            fail_tag_column = headers.index("fail_tag") + 1
            fail_tag_cell = next(
                row[fail_tag_column - 1]
                for row in sheet.iter_rows(min_row=2)
                if row[headers.index("component")].value == "takeaway"
                and row[headers.index("component_index")].value == 1
            )
            self.assertEqual(fail_tag_cell.fill.fgColor.rgb[-6:], FAIL_TAG_COLORS["wordcount"])

    def test_takeaway_review_uses_single_text_not_title_description_pair(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][3]["content"]["takeaways"][0] = {
                "title": (
                    "Energy access progress depends on affordability, reliability, productive uses, and local service quality, "
                    "not only whether a household or institution is physically connected to infrastructure."
                ),
                "description": "",
            }
            save_json(lesson, module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            takeaway_rows = [
                row for row in output.component_rows
                if row["component"] == "takeaway" and row["component_index"] == 1
            ]
            self.assertEqual({row["criterion"] for row in takeaway_rows}, {"has takeaway text", "takeaway text word count 6-45"})
            self.assertTrue(all(row["status"] == "pass" for row in takeaway_rows))

    def test_takeaway_word_count_uses_title_and_description_together(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][3]["content"]["takeaways"][0] = {
                "title": "Energy access progress depends on affordable and reliable service.",
                "description": (
                    "Learners should connect infrastructure planning with service quality, local institutions, productive uses, "
                    "household affordability, community priorities, long-term maintenance, and the practical conditions that "
                    "determine whether electricity access improves daily life, strengthens livelihoods, and supports resilient "
                    "public services over time while reducing avoidable disruption for schools, clinics, farms, and small "
                    "businesses."
                ),
            }
            save_json(lesson, module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            word_count_row = next(
                row for row in output.component_rows
                if row["component"] == "takeaway"
                and row["component_index"] == 1
                and row["criterion"] == "takeaway text word count 6-45"
            )
            self.assertEqual(word_count_row["status"], "fail")
            self.assertIn("Reduce takeaway to 45 words maximum", word_count_row["notes"])
            self.assertIn("Energy access progress", word_count_row["value"])
            self.assertIn("Learners should connect", word_count_row["value"])

    def test_updated_word_count_controls_are_applied(self) -> None:
        controls = load_review_controls()
        self.assertEqual(controls["takeaways"]["text_max_words"], 45)
        self.assertEqual(controls["concepts"]["title_max_words"], 6)
        self.assertEqual(controls["concepts"]["definition_max_words"], 32)
        self.assertEqual(controls["case_study"]["subtitle_max_words"], 20)

    def test_word_count_max_uses_buffer_before_failing(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][0]["content"]["concepts"][0]["title"] = "Seven word concept title should still pass"
            save_json(lesson, module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            title_rows = [
                row for row in output.component_rows
                if row["component"] == "concept"
                and row["component_index"] == 1
                and row["criterion"] == "title word count 1-6"
            ]
            self.assertEqual(len(title_rows), 1)
            self.assertEqual(title_rows[0]["status"], "pass")

    def test_quiz_and_poll_word_count_rules_are_applied(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][2]["content"]["title"] = "Vote?"
            lesson["segments"][2]["content"]["options"] = [
                {
                    "id": 1,
                    "value": (
                        "This answer option is intentionally much too long for a poll response option and should fail "
                        "because it keeps adding unnecessary extra wording"
                    ),
                },
                {"id": 2, "value": "Reliability"},
            ]
            save_json(lesson, module_dir / "1.1.1.json")
            save_json(
                {
                    "id": "1.-1.0",
                    "segments": [
                        {
                            "template_id": "scored_quiz",
                            "content": {
                                "questions": [
                                    {
                                        "id": 1,
                                        "prompt": "Why?",
                                        "options": [
                                            {"id": 1, "value": "A"},
                                            {
                                                "id": 2,
                                                "value": (
                                                    "This quiz answer option is intentionally much too long for the configured "
                                                    "answer option limit and should definitely fail the review"
                                                ),
                                            },
                                            {"id": 3, "value": "C"},
                                        ],
                                        "solution": [1],
                                        "messages": {"correct": {"body": "Yes"}, "wrong": {"body": "No"}},
                                    }
                                ]
                            },
                        }
                    ],
                },
                module_dir / "1.-1.0.json",
            )
            output = run_deterministic_review(load_review_inputs(root))
            poll_rows = [row for row in output.component_rows if row["component"] == "poll"]
            quiz_rows = [row for row in output.component_rows if row["component"] == "quiz_question"]
            self.assertTrue(any(row["criterion"] == "question word count 5-30" and row["status"] == "fail" for row in poll_rows))
            self.assertTrue(any(row["criterion"] == "option word count 1-15" and row["status"] == "fail" for row in poll_rows))
            self.assertTrue(any("Reduce poll option" in row["notes"] for row in poll_rows))
            self.assertTrue(any(row["criterion"] == "prompt word count 5-35" and row["status"] == "fail" for row in quiz_rows))
            self.assertTrue(any(row["criterion"] == "option word count 1-18" and row["status"] == "fail" for row in quiz_rows))
            self.assertTrue(any("Reduce quiz option" in row["notes"] for row in quiz_rows))

    def test_case_study_checks_subtitle_max_but_not_body_after_cover(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"].append(
                {
                    "template_id": "case_study_cover",
                    "content": {
                        "title": "Mini-grid planning",
                        "subtitle": (
                            "This subtitle is intentionally much too long for the configured case study review limit "
                            "and includes several extra words to ensure the updated twenty word maximum still fails"
                        ),
                    },
                }
            )
            save_json(lesson, module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            case_rows = [row for row in output.component_rows if row["component"] == "case_study"]
            self.assertFalse(any(row["criterion"] == "has following content" for row in case_rows))
            self.assertTrue(any(row["criterion"] == "subtitle word count 0-20" and row["status"] == "fail" for row in case_rows))
            self.assertTrue(any("Reduce case study subtitle" in row["notes"] for row in case_rows))

    def test_llm_payloads_are_indexed_per_item(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            lesson = load_review_inputs(root)[0]
            payload = _component_payload(lesson, "takeaways")
            self.assertEqual(payload[0]["index"], 1)
            self.assertIn("takeaway_text", payload[0])
            self.assertNotIn("title", payload[0])
            self.assertNotIn("description", payload[0])
            self.assertIn("Energy access requires more than grid connections.", payload[0]["takeaway_text"])
            self.assertIn("Useful access depends on affordability", payload[0]["takeaway_text"])

    def test_llm_request_tells_model_to_ignore_formatting_artifacts(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][3]["content"]["takeaways"][0]["description"] = (
                "- Useful access depends on affordability, reliability, safety, and services."
            )
            save_json(lesson, module_dir / "1.1.1.json")
            review_input = load_review_inputs(root)[0]
            request = _build_llm_review_request(
                review_input,
                "takeaways",
                _component_payload(review_input, "takeaways"),
                load_review_controls(),
                review_input.lesson_text,
            )

            self.assertIn("content_sections", request["lesson_context"])
            self.assertIn("line breaks, bullets", request["review_instructions"])
            self.assertIn("hyphen separators", request["review_instructions"])
            self.assertNotIn("component_payload", request)
            self.assertNotIn("review_context", request)

    def test_concept_llm_request_includes_definition_and_lesson_context(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            save_json(
                {
                    "modules": [
                        {
                            "id": "1",
                            "title": "Energy Access",
                            "chapters": [
                                {
                                    "id": "1.1",
                                    "title": "Access Foundations",
                                    "lessons": [
                                        {"id": "1.1.1", "title": "Lesson 1: <strong>Energy Access Basics</strong>"}
                                    ],
                                }
                            ],
                        }
                    ]
                },
                root / "module_structure.json",
            )
            lesson = load_review_inputs(root)[0]
            payload = _component_payload(lesson, "concepts")
            request = _build_llm_review_request(
                lesson,
                "concepts",
                payload,
                load_review_controls(),
                lesson.lesson_text,
            )
            self.assertEqual(payload[0]["concept"], "Energy access")
            self.assertIn("definition", payload[0])
            self.assertEqual(request["lesson_context"]["lesson_title"], "Lesson 1: Energy Access Basics")
            self.assertIn("This lesson explains energy access", request["lesson_context"]["content_excerpt"])
            self.assertIn("For concepts", request["review_instructions"])
            self.assertIn("glossary", request["allowed_fail_tags"])
            self.assertIn("fail_tag", request["response_schema"]["items"][0])

    def test_llm_parser_accepts_per_item_results(self) -> None:
        parsed = _parse_llm_json(
            '{"passed": false, "notes": "one issue", "items": '
            '[{"index": 2, "passed": false, "fail_tag": "off_topic", "notes": "off topic", '
            '"suggested_refinement": "Tie this concept to energy access planning."}]}'
        )
        self.assertFalse(parsed["passed"])
        self.assertEqual(parsed["items"][0]["index"], 2)
        self.assertFalse(parsed["items"][0]["passed"])
        self.assertEqual(parsed["items"][0]["fail_tag"], "offtopic")
        self.assertEqual(parsed["items"][0]["suggested_refinement"], "Tie this concept to energy access planning.")

    def test_llm_parser_uses_item_results_over_contradictory_group_status(self) -> None:
        parsed = _parse_llm_json(
            '{"passed": false, "notes": "generic group warning", "items": '
            '[{"index": 1, "passed": true, "fail_tag": "none", "notes": "", "suggested_refinement": ""}]}'
        )

        self.assertTrue(parsed["passed"])

    def test_resource_llm_request_does_not_include_lesson_body_context(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            review_input = load_review_inputs(root)[0]
            request = _build_llm_review_request(
                review_input,
                "resources",
                _component_payload(review_input, "resources"),
                load_review_controls(),
                review_input.lesson_text,
            )

            self.assertEqual(request["lesson_context"]["content_sections"], [])
            self.assertEqual(request["lesson_context"]["content_excerpt"], "")
            self.assertIn("resource_review_context", request["lesson_context"])
            self.assertIn("those are not resource quality issues", request["review_instructions"])

    def test_critical_llm_request_omits_lesson_body_context(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            review_input = load_review_inputs(root)[0]
            request = _build_llm_review_request(
                review_input,
                "takeaways",
                _component_payload(review_input, "takeaways"),
                load_review_controls(),
                review_input.lesson_text,
                review_mode="critical",
            )

            self.assertEqual(request["lesson_context"]["content_sections"], [])
            self.assertEqual(request["lesson_context"]["content_excerpt"], "")
            self.assertIn("critical_review_context", request["lesson_context"])

    def test_component_detail_workbook_keeps_llm_fail_tag_and_fix(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            output.component_rows.append(
                {
                    "lesson_id": "1.1.1",
                    "module_id": "1",
                    "chapter_id": "1.1",
                    "lesson_type": "lesson",
                    "component": "concept",
                    "component_index": 1,
                    "criterion": "llm semantic review",
                    "status": "fail",
                    "fail_tag": "unclear_definition",
                    "notes": "The definition is too broad for the lesson context.",
                    "suggested_fix": "Define how energy access applies to households and services in this lesson.",
                    "value": "Energy access",
                    "source_json": str(module_dir / "1.1.1.json"),
                }
            )
            workbook_path = write_review_workbook(output, root / "review.xlsx", keep_pass=False)
            workbook = load_workbook(workbook_path)
            sheet = workbook["Component Details"]
            headers = [cell.value for cell in sheet[1]]
            self.assertIn("context_accurate_ai_review", headers)
            rows = [
                {headers[index]: value for index, value in enumerate(row)}
                for row in sheet.iter_rows(min_row=2, values_only=True)
            ]
            concept_row = next(row for row in rows if row["component"] == "concept" and row["component_index"] == 1)
            self.assertEqual(concept_row["fail_tag"], "unclear")
            self.assertIn("Fix: Define how energy access applies", concept_row["summary"])
            self.assertIn("Concept: Energy access", concept_row["value"])
            self.assertIn("Definition:", concept_row["value"])
            fail_tag_column = headers.index("fail_tag") + 1
            fail_tag_cell = next(
                row[fail_tag_column - 1]
                for row in sheet.iter_rows(min_row=2)
                if row[headers.index("component")].value == "concept"
                and row[headers.index("component_index")].value == 1
            )
            self.assertEqual(fail_tag_cell.fill.fgColor.rgb[-6:], FAIL_TAG_COLORS["unclear"])

    def test_concept_definition_quality_not_checked_by_naive_word_match(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            lesson = _normal_lesson()
            lesson["segments"][0]["content"]["concepts"][0] = {
                "title": "Just transition",
                "body": "A fair shift toward cleaner energy systems that protects workers, communities, and vulnerable groups while expanding opportunity.",
            }
            save_json(lesson, module_dir / "1.1.1.json")
            output = run_deterministic_review(load_review_inputs(root))
            concept_rows = [
                row for row in output.component_rows
                if row["component"] == "concept" and row["component_index"] == 1
            ]
            self.assertFalse(any(row["criterion"] == "definition references concept" for row in concept_rows))
            self.assertTrue(all(row["status"] == "pass" for row in concept_rows))

    def test_quiz_llm_context_uses_module_content_when_available(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            module_dir = root / "Module_1"
            save_json(_normal_lesson("1.1.1"), module_dir / "1.1.1.json")
            save_json(
                {
                    "id": "1.-1.0",
                    "segments": [
                        {
                            "template_id": "scored_quiz",
                            "content": {
                                "questions": [
                                    {
                                        "id": 1,
                                        "prompt": "What does energy access require?",
                                        "options": [{"id": 1, "value": "Affordability"}, {"id": 2, "value": "Nothing"}],
                                        "solution": [1],
                                        "messages": {"correct": {"body": "Yes"}, "wrong": {"body": "No"}},
                                    }
                                ]
                            },
                        }
                    ],
                },
                module_dir / "1.-1.0.json",
            )
            lessons = load_review_inputs(root)
            contexts = _module_contexts(lessons, load_review_controls())
            self.assertIn("This lesson explains energy access", contexts["1"])

    def test_component_session_accumulates_and_writes_workbook(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            lesson_root = root / "03_Outputs" / "SEA_Modules" / "en"
            module_dir = lesson_root / "Module_1"
            save_json(_normal_lesson(), module_dir / "1.1.1.json")
            config = PipelineConfig(
                root_dir=root,
                paths=RepoPaths.from_root(root),
                env=EnvConfig(),
                stage_options={
                    "review_lesson_root": lesson_root,
                    "review_test_start_lesson_id": "1.1.1",
                    "review_test_lesson_count": 1,
                    "review_output_path": root / "review.xlsx",
                    "review_cache_path": root / "cache.json",
                    "review_use_llm": False,
                    "review_check_resource_blobs": False,
                },
            )
            session = create_review_session(config, log=None)
            run_review_component(session, "takeaways", log=None)
            self.assertEqual(session.output.lesson_rows[0]["takeaways_status"], "pass")
            self.assertEqual(session.output.lesson_rows[0]["concepts_status"], "not_run")
            run_review_component(session, "concepts", log=None)
            result = write_review_session_workbook(session, log=None)
            self.assertTrue(Path(result["summary"]["output_path"]).exists())
            workbook = load_workbook(result["summary"]["output_path"])
            headers = [cell.value for cell in workbook["Component Details"][1]]
            self.assertNotIn("context_accurate_ai_review", headers)


if __name__ == "__main__":
    unittest.main()
