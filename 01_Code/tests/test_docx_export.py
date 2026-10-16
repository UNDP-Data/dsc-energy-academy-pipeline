from __future__ import annotations

import base64
import csv
import json
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

CODE_DIR = Path(__file__).resolve().parents[1]
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from sea_pipeline.auxiliary.docx_export import export_docx_materials
from sea_pipeline.auxiliary.runner import run_auxiliary_stage
from sea_pipeline.shared.config import PipelineConfig
from sea_pipeline.shared.env import EnvConfig
from sea_pipeline.shared.io_json import save_json
from sea_pipeline.shared.paths import RepoPaths


PNG_1X1 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def _config(root: Path, **stage_options: object) -> PipelineConfig:
    return PipelineConfig(
        root_dir=root,
        paths=RepoPaths.from_root(root),
        env=EnvConfig(),
        stage_options=stage_options,
    )


def _write_fixture(root: Path) -> Path:
    lesson_root = root / "03_Outputs" / "SEA_Modules" / "en"
    module_dir = lesson_root / "Module_1"
    image_path = root / "02_Inputs" / "fixture_infographic.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_path.write_bytes(base64.b64decode(PNG_1X1))
    save_json(
        {
            "modules": [
                {
                    "id": "1",
                    "title": "Foundations",
                    "description": "Module description.",
                    "chapters": [
                        {
                            "id": "1.1",
                            "title": "Chapter <strong>One</strong>",
                            "lessons": [
                                {"id": "1.1.1", "title": "Lesson <strong>One</strong>"},
                            ],
                        }
                    ],
                }
            ]
        },
        lesson_root / "module_structure.json",
    )
    save_json(
        {
            "id": "1.1.1",
            "segments": [
                {
                    "template_id": "text",
                    "content": {
                        "text_elements": [
                            {"template_id": "subtitle", "content": {"text": "Large subtitle text"}},
                            {"template_id": "paragraph_large", "content": {"text": "Large paragraph text"}},
                            {
                                "template_id": "paragraph_medium",
                                "content": {
                                    "text": "Main <strong>lesson</strong> text with <ai prompt=\"x\">AI phrase</ai> and <a href=\"https://example.org/source\"><u>source link</u></a>."
                                },
                            },
                            {"template_id": "paragraph_small", "content": {"text": "Small paragraph text"}},
                            {"template_id": "kpi_highlight_medium", "content": {"text": "KPI highlight text"}},
                            {"template_id": "bullet_point", "content": {"text": "A visible bullet."}},
                        ]
                    },
                },
                {
                    "template_id": "list_of_lessons",
                    "content": {
                        "title": "What comes next?",
                        "lessons": [{"lessonId": "1.1.2", "title": "Next lesson", "description": "Next lesson description"}],
                    },
                },
                {
                    "template_id": "learning_objectives",
                    "content": {
                        "title": "Learning Objectives",
                        "objectives": [{"title": "Explain energy access.", "description": "Connect it to development."}],
                    },
                },
                {
                    "template_id": "key_resources",
                    "content": {
                        "title": "Key Resources",
                        "resources": [{"text": "Resource title", "href": "https://example.org/resource"}],
                    },
                },
                {
                    "template_id": "infographic",
                    "content": {
                        "image": {"src": image_path.as_uri(), "caption": "Infographic caption"},
                        "metadata": {
                            "title": "Infographic title",
                            "subtitle": "Infographic subtitle",
                            "description": "Infographic description",
                            "footnote": "Infographic source note",
                        },
                    },
                },
                {
                    "template_id": "echarts_chart",
                    "content": {
                        "chart_id": "M1_C1_1",
                        "title": "Chart title",
                        "metadata": {"footnote": "Chart footnote"},
                    },
                },
                {
                    "template_id": "embed",
                    "content": {"title": "Interactive map", "url": "https://example.org/embed"},
                },
                {
                    "template_id": "video",
                    "content": {
                        "src": "https://sehseadata.blob.core.windows.net/images/Avatars/en/1/Videos/Lesson-1.4.0_en.mp4",
                        "poster": "https://sehseadata.blob.core.windows.net/images/Avatars/en/1/Photos/Lesson-1.4.0_en.webp",
                    },
                },
                {
                    "template_id": "poll",
                    "content": {
                        "title": "Poll question",
                        "description": "Poll description text",
                        "options": [{"id": 1, "value": "First option"}],
                    },
                },
                {
                    "template_id": "image-hotspot",
                    "content": {
                        "image": {"src": "https://example.org/hotspot.webp", "caption": ""},
                        "hotspots": [{"label": "Hotspot label", "description": "Hotspot body"}],
                    },
                },
                {
                    "template_id": "scored_quiz",
                    "content": {
                        "title": "Quiz",
                        "intro": "Quiz intro",
                        "passingScore": 0.8,
                        "questions": [
                            {
                                "id": 1,
                                "prompt": "What does energy access support?",
                                "options": [{"id": 1, "value": "Development"}, {"id": 2, "value": "Nothing"}],
                                "solution": [1],
                                "messages": {"correct": {"body": "Correct feedback"}, "wrong": {"body": "Wrong feedback"}},
                            }
                        ],
                    },
                },
            ],
        },
        module_dir / "1.1.1.json",
    )
    return lesson_root


class DocxExportTests(unittest.TestCase):
    def test_exports_complete_material_docx_with_images_and_placeholders(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            lesson_root = _write_fixture(root)
            config = _config(root, docx_export_lesson_root=lesson_root)
            result = export_docx_materials(config, with_images=True, log=None)
            self.assertEqual(result["status"], "ok")
            output_path = Path(result["outputs"][0])
            document = Document(output_path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Main lesson text with AI phrase and source link (https://example.org/source).", text)
            self.assertIn("A visible bullet.", text)
            self.assertIn("1.1.2 | Next lesson | Next lesson description", text)
            self.assertIn("Explain energy access. Connect it to development.", text)
            self.assertIn("Resource title (https://example.org/resource)", text)
            self.assertIn("Infographic title", text)
            self.assertIn("Infographic source note", text)
            self.assertIn("[CHART PLACEHOLDER: M1_C1_1 | Chart title]", text)
            self.assertIn("[EMBED PLACEHOLDER: Interactive map | https://example.org/embed]", text)
            self.assertIn("Video: Avatars/en/1/Videos/Lesson-1.4.0_en.mp4", text)
            self.assertNotIn("[VIDEO PLACEHOLDER:", text)
            self.assertNotIn("Poster:", text)
            self.assertNotIn("https://sehseadata.blob.core.windows.net/images/Avatars/en/1/Videos/Lesson-1.4.0_en.mp4", text)
            self.assertIn("Poll description text", text)
            self.assertIn("[IMAGE HOTSPOT PLACEHOLDER: https://example.org/hotspot.webp | hotspots: 1]", text)
            self.assertIn("Hotspot 1: Hotspot label | Hotspot body", text)
            self.assertIn("What does energy access support?", text)
            self.assertIn("Correct feedback", text)
            self.assertEqual(len(document.inline_shapes), 1)
            sizes = {
                paragraph.text: paragraph.runs[0].font.size.pt
                for paragraph in document.paragraphs
                if paragraph.text in {"Large subtitle text", "Large paragraph text", "Small paragraph text", "KPI highlight text", "A visible bullet."}
                and paragraph.runs
                and paragraph.runs[0].font.size
            }
            self.assertEqual(sizes["Large subtitle text"], 16)
            self.assertEqual(sizes["Large paragraph text"], 13)
            self.assertEqual(sizes["Small paragraph text"], 9)
            self.assertEqual(sizes["KPI highlight text"], 15)
            self.assertEqual(sizes["A visible bullet."], 11)

    def test_docx_export_stage_uses_module_implementation(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            lesson_root = _write_fixture(root)
            config = _config(root, docx_export_lesson_root=lesson_root, docx_export_module_ids="1")
            result = run_auxiliary_stage(config, "docx_export_text")
            self.assertEqual(result["kind"], "python")
            self.assertEqual(result["summary"]["modules_exported"], 1)
            self.assertEqual(result["summary"]["images_inserted"], 0)
            self.assertTrue(Path(result["outputs"][0]).exists())

    def test_failed_infographic_image_is_audited_and_placeholdered(self) -> None:
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            lesson_root = _write_fixture(root)
            lesson_path = lesson_root / "Module_1" / "1.1.1.json"
            lesson = json.loads(lesson_path.read_text(encoding="utf-8"))
            lesson["segments"][4]["content"]["image"]["src"] = "file:///missing/image.webp"
            save_json(lesson, lesson_path)
            config = _config(root, docx_export_lesson_root=lesson_root)
            result = export_docx_materials(config, with_images=True, log=None)
            self.assertEqual(result["status"], "warning")
            document = Document(Path(result["outputs"][0]))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("[INFOGRAPHIC IMAGE: file:///missing/image.webp]", text)
            with open(result["summary"]["audit_path"], encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertTrue(any(row["issue_type"] == "image_download_failed" for row in rows))


if __name__ == "__main__":
    unittest.main()
